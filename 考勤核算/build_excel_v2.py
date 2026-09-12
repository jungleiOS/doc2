# -*- coding: utf-8 -*-
"""生成「活公式」版8月统计 Excel（v3，分钟取整避免浮点误差）：
- 每日明细：首卡/末卡为数据，状态/跨度分钟/跨度h/取整/白班/加班全是公式
- 汇总表：COUNTIFS/SUMIFS 引用明细，改明细自动重算
- 口径说明 sheet"""
import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/jiangjun/Documents/Kimi/Workspaces/zhihu/考勤核算"
with open(f"{BASE}/august_result.json", encoding="utf-8") as f:
    data = json.load(f)
people, summary, order = data["people"], data["summary"], data["order"]

wb = Workbook()
thin = Border(*[Side(style="thin")]*4)
center = Alignment(horizontal="center", vertical="center")
yellow = PatternFill("solid", fgColor="FFF2CC")
redfont = Font(color="CC0000")
bold = Font(bold=True)
NROW = 1 + 31 * len(order)
LAST = str(NROW)
DET = "'每日明细'"

# ---------- sheet 每日明细 ----------
ws2 = wb.active
ws2.title = "每日明细"
heads = ["姓名","日期","星期","当日打卡","首卡","末卡","状态","跨度(分钟)","跨度(h)","取整工时","白班","加班"]
for j, h in enumerate(heads, 1):
    c = ws2.cell(row=1, column=j, value=h)
    c.font = bold; c.alignment = center; c.border = thin
wk = ["一","二","三","四","五","六","日"]
r = 2
for name in order:
    days = people[name]
    for d in range(1, 32):
        info = days[str(d)]
        date = datetime.date(2026, 8, d)
        ws2.cell(row=r, column=1, value=name)
        ws2.cell(row=r, column=2, value=f"8/{d}")
        ws2.cell(row=r, column=3, value=wk[date.weekday()])
        ws2.cell(row=r, column=4, value=" ".join(info["times"]))
        t = info["times"]
        if len(t) >= 2:
            h1, m1 = map(int, t[0].split(":")); h2, m2 = map(int, t[-1].split(":"))
            ws2.cell(row=r, column=5, value=datetime.time(h1, m1))
            ws2.cell(row=r, column=6, value=datetime.time(h2, m2))
        elif len(t) == 1:
            h1, m1 = map(int, t[0].split(":"))
            ws2.cell(row=r, column=5, value=datetime.time(h1, m1))
        ws2.cell(row=r, column=7,  value=f'=IF($E{r}="","空卡",IF($F{r}="","单卡","正常"))')
        ws2.cell(row=r, column=8,  value=f'=IF($G{r}="正常",ROUND(($F{r}-$E{r})*1440,0),"")')
        ws2.cell(row=r, column=9,  value=f'=IF($G{r}="正常",$H{r}/60,"")')
        ws2.cell(row=r, column=10, value=f'=IF($G{r}="正常",ROUND($H{r}/30,0)/2,"")')
        ws2.cell(row=r, column=11, value=f'=IF($G{r}="正常",MIN($J{r},12),"")')
        ws2.cell(row=r, column=12, value=f'=IF($G{r}="正常",MAX($J{r}-12,0),"")')
        for j in range(1, 13):
            c = ws2.cell(row=r, column=j)
            c.border = thin; c.alignment = center
            if info["status"] != "正常":
                c.fill = yellow
                if j == 7: c.font = redfont
        ws2.cell(row=r, column=5).number_format = "hh:mm"
        ws2.cell(row=r, column=6).number_format = "hh:mm"
        ws2.cell(row=r, column=8).number_format = "0"
        ws2.cell(row=r, column=9).number_format = "0.00"
        for j in (10, 11, 12):
            ws2.cell(row=r, column=j).number_format = "0.0"
        r += 1
for w, col in zip([10,7,6,22,8,8,7,10,9,9,8,8], range(1,13)):
    ws2.column_dimensions[get_column_letter(col)].width = w
ws2.freeze_panes = "A2"

# ---------- sheet 汇总 ----------
ws = wb.create_sheet("8月白班及加班统计", 0)
ws["A1"] = "2026年8月考勤明细表（白班封顶12h/天，超出计加班；明细补卡后本表自动重算）"
ws["A1"].font = Font(bold=True, size=13)
ws.merge_cells("A1:K1")
head = ["姓名","上班天数","白班","加班","合计"]
for i, h in enumerate(head):
    for col0 in (1, 7):
        c = ws.cell(row=2, column=col0+i, value=h)
        c.font = bold; c.alignment = center; c.border = thin

half = 31
left, right = order[:half], order[half:]
for i in range(half):
    r = 3 + i
    ws.cell(row=r, column=1, value=left[i])
    cnt = f'COUNTIFS({DET}!$A$2:$A${LAST},$A{r},{DET}!$G$2:$G${LAST},"正常")'
    fs = [f'=IF({cnt}=0,"/",{cnt})',
          f'=IF($B{r}="/","/",SUMIFS({DET}!$K$2:$K${LAST},{DET}!$A$2:$A${LAST},$A{r}))',
          f'=IF($B{r}="/","/",SUMIFS({DET}!$L$2:$L${LAST},{DET}!$A$2:$A${LAST},$A{r}))',
          f'=IF($B{r}="/","/",$C{r}+$D{r})']
    for j, f in enumerate(fs):
        ws.cell(row=r, column=2+j, value=f)
for i in range(len(right)):
    r = 3 + i
    ws.cell(row=r, column=7, value=right[i])
    cnt = f'COUNTIFS({DET}!$A$2:$A${LAST},$G{r},{DET}!$G$2:$G${LAST},"正常")'
    fs = [f'=IF({cnt}=0,"/",{cnt})',
          f'=IF($H{r}="/","/",SUMIFS({DET}!$K$2:$K${LAST},{DET}!$A$2:$A${LAST},$G{r}))',
          f'=IF($H{r}="/","/",SUMIFS({DET}!$L$2:$L${LAST},{DET}!$A$2:$A${LAST},$G{r}))',
          f'=IF($H{r}="/","/",$I{r}+$J{r})']
    for j, f in enumerate(fs):
        ws.cell(row=r, column=8+j, value=f)
for r in range(3, 3+half):
    for c in list(range(1,6)) + list(range(7,12)):
        cell = ws.cell(row=r, column=c)
        cell.border = thin; cell.alignment = center
for w, col in zip([10,10,9,9,9,2,10,10,9,9,9], range(1,12)):
    ws.column_dimensions[get_column_letter(col)].width = w

# ---------- sheet 口径说明 ----------
ws3 = wb.create_sheet("口径说明")
lines = [
    ("2026年8月考勤核算口径与公式说明", True),
    ("", False),
    ("一、核算口径（已用7月数据验证：45人中白班40人、加班41人完全吻合）", True),
    ("1. 上班天数：当日首末卡齐全（状态=正常）记 1 天；单卡、空卡不计，留白待人工补卡", False),
    ("2. 当日跨度(分钟) = ROUND((末卡 − 首卡) × 1440, 0) —— 先化成整数分钟再算，避免浮点误差", False),
    ("3. 取整工时 = ROUND(跨度分钟 ÷ 30, 0) ÷ 2 —— 四舍五入到 0.5 小时（30分钟为一个单位）", False),
    ("4. 白班 = MIN(取整工时, 12)：每天最多计 12 小时白班", False),
    ("5. 加班 = MAX(取整工时 − 12, 0)：超出 12 小时的部分才算加班", False),
    ("6. 月合计 = 白班合计 + 加班合计", False),
    ("", False),
    ("二、每日明细公式（第2行为例，向下填充）", True),
    ("状态      G2 = IF($E2=\"\",\"空卡\",IF($F2=\"\",\"单卡\",\"正常\"))", False),
    ("跨度(分)  H2 = IF($G2=\"正常\",ROUND(($F2-$E2)*1440,0),\"\")", False),
    ("跨度(h)   I2 = IF($G2=\"正常\",$H2/60,\"\")", False),
    ("取整工时  J2 = IF($G2=\"正常\",ROUND($H2/30,0)/2,\"\")", False),
    ("白班      K2 = IF($G2=\"正常\",MIN($J2,12),\"\")", False),
    ("加班      L2 = IF($G2=\"正常\",MAX($J2-12,0),\"\")", False),
    ("", False),
    ("三、汇总表公式（以左栏第3行为例）", True),
    ("上班天数 B3 = IF(COUNTIFS(明细!$A$2:$A$1892,$A3,明细!$G$2:$G$1892,\"正常\")=0,\"/\",COUNTIFS(...))", False),
    ("白班 C3 = IF($B3=\"/\",\"/\",SUMIFS(明细!$K$2:$K$1892,明细!$A$2:$A$1892,$A3))", False),
    ("加班 D3 = IF($B3=\"/\",\"/\",SUMIFS(明细!$L$2:$L$1892,明细!$A$2:$A$1892,$A3))", False),
    ("合计 E3 = IF($B3=\"/\",\"/\",$C3+$D3)", False),
    ("", False),
    ("四、使用方法", True),
    ("1. 在「每日明细」找到标黄的行（单卡/空卡），补上缺失的首卡或末卡（格式 hh:mm）", False),
    ("2. 状态自动变「正常」，跨度/取整/白班/加班自动算出", False),
    ("3. 「8月白班及加班统计」用 COUNTIFS/SUMIFS 引用明细，自动重算，无需手改", False),
    ("4. 如口径调整（例如封顶改 11.5h），只需改明细 K、L 两列公式再向下填充", False),
    ("", False),
    ("五、为什么先算分钟再算小时？", True),
    ("时间相减会产生二进制浮点误差：如 44100秒÷86400×24 在计算机里 = 12.249999999999998。", False),
    ("若直接 ROUND(×24×2) 会把本应是 12.5 的值判成 12.0（差 0.5 小时）。", False),
    ("先 ROUND(×1440) 得到整数分钟，整数÷30 在边界处（如 735÷30=24.5）是精确值， ROUND 才稳定。", False),
]
for i, (txt, b) in enumerate(lines, 1):
    c = ws3.cell(row=i, column=1, value=txt)
    if b: c.font = bold
ws3.column_dimensions["A"].width = 105

out = f"{BASE}/08月白班及加班统计.xlsx"
wb.save(out)
print("saved:", out)
