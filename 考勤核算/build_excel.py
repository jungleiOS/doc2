# -*- coding: utf-8 -*-
"""生成8月白班及加班统计 Excel：sheet1 汇总（仿7月版式），sheet2 每日明细（异常天标黄留白）。"""
import json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = "/Users/jiangjun/Documents/Kimi/Workspaces/zhihu/考勤核算"
with open(f"{BASE}/august_result.json", encoding="utf-8") as f:
    data = json.load(f)
people, summary, order = data["people"], data["summary"], data["order"]

wb = Workbook()
thin = Border(*[Side(style="thin")]*4)
center = Alignment(horizontal="center", vertical="center")
yellow = PatternFill("solid", fgColor="FFF2CC")
redfont = Font(color="CC0000")

# ---------- sheet1 汇总 ----------
ws = wb.active
ws.title = "8月白班及加班统计"
ws["A1"] = "2026年8月考勤明细表（按打卡核算：白班封顶12h/天，超出计加班）"
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:K1")
head = ["姓名", "上班天数", "白班", "加班", "合计"]
for i, h in enumerate(head):
    c = ws.cell(row=2, column=1+i, value=h)
    c.font = Font(bold=True); c.alignment = center; c.border = thin
    c2 = ws.cell(row=2, column=7+i, value=h)
    c2.font = Font(bold=True); c2.alignment = center; c2.border = thin

half = 31
left, right = order[:half], order[half:]
def fmt(person):
    sm = summary[person]
    if sm["days"] == 0:
        return [person, "/", "/", "/", "/"]
    return [person, sm["days"], sm["bai"], sm["jia"], round(sm["bai"]+sm["jia"], 1)]

for i in range(half):
    r = 3 + i
    for j, v in enumerate(fmt(left[i])):
        c = ws.cell(row=r, column=1+j, value=v); c.border = thin; c.alignment = center
for i in range(len(right)):
    r = 3 + i
    for j, v in enumerate(fmt(right[i])):
        c = ws.cell(row=r, column=7+j, value=v); c.border = thin; c.alignment = center
for w, col in zip([10,10,9,9,9,2,10,10,9,9,9], range(1,12)):
    ws.column_dimensions[get_column_letter(col)].width = w

# ---------- sheet2 每日明细 ----------
ws2 = wb.create_sheet("每日明细")
heads = ["姓名","日期","星期","当日打卡","首卡","末卡","跨度(h)","取整工时","白班","加班","状态"]
for j, h in enumerate(heads, 1):
    c = ws2.cell(row=1, column=j, value=h)
    c.font = Font(bold=True); c.alignment = center; c.border = thin
wk = ["一","二","三","四","五","六","日"]
r = 2
for name in order:
    days = people[name]
    for d in range(1, 32):
        info = days[str(d)] if str(d) in days else days[d]
        date = datetime.date(2026, 8, d)
        punch = " ".join(info["times"])
        row = [name, f"8/{d}", wk[date.weekday()], punch,
               info["times"][0] if info["times"] else "",
               info["times"][-1] if info["times"] else "",
               info["span"] if info["span"] is not None else "",
               info["rounded"] if info["rounded"] is not None else "",
               info["bai"] if info["bai"] is not None else "",
               info["jia"] if info["jia"] is not None else "",
               info["status"]]
        for j, v in enumerate(row, 1):
            c = ws2.cell(row=r, column=j, value=v)
            c.border = thin; c.alignment = center
            if info["status"] != "正常":
                c.fill = yellow
                if j == 11: c.font = redfont
        r += 1
for w, col in zip([10,7,6,22,7,7,9,9,8,8,7], range(1,12)):
    ws2.column_dimensions[get_column_letter(col)].width = w
ws2.freeze_panes = "A2"

out = f"{BASE}/08月白班及加班统计.xlsx"
wb.save(out)
print("saved:", out)
